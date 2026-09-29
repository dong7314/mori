"""Small document worker. Fixed workspace, new outputs, no network or code execution.

Not a tenant sandbox: only enable for a trusted single-user Hermes profile.
"""

import io
import json
import math
import os
import re
import sys
import uuid
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

MAX_FILE = 10_000_000
MAX_TEXT = 40_000
FONT = Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf")


def fields(obj, allowed):
    if not isinstance(obj, dict) or set(obj) - set(allowed):
        raise ValueError("Unsupported fields")


def text(value):
    if not isinstance(value, str) or len(value) > MAX_TEXT:
        raise ValueError("Expected text of at most 40000 characters")
    if any(ord(c) < 32 and c not in "\n\t\r" for c in value):
        raise ValueError("Control characters are not supported")
    return value


def items(value, maximum=200):
    if not isinstance(value, list) or len(value) > maximum:
        raise ValueError(f"Expected a list with at most {maximum} items")
    return value


def index(value, maximum):
    if type(value) is not int or not 0 <= value < maximum:
        raise ValueError("Index out of range")
    return value


def cell_value(value):
    if value is None or type(value) is bool:
        return value
    if type(value) in (int, float) and math.isfinite(value) and abs(value) <= 1e100:
        return value
    if isinstance(value, str) and len(value) <= 32767:
        return text(value)
    raise ValueError("Invalid spreadsheet value")


def read_input(root, file_id):
    if not isinstance(file_id, str) or not re.fullmatch(r"[0-9a-f]{32}\.(docx|xlsx|pdf)", file_id):
        raise ValueError("Invalid file_id")
    fd = os.open(root / file_id, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
    with os.fdopen(fd, "rb") as stream:
        import stat

        if not stat.S_ISREG(os.fstat(stream.fileno()).st_mode):
            raise ValueError("Expected a regular file")
        data = stream.read(MAX_FILE + 1)
    if len(data) > MAX_FILE:
        raise ValueError("Input exceeds 10 MB")
    fmt = file_id.rsplit(".", 1)[1]
    if fmt != "pdf":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 50_000_000:
                raise ValueError("Office archive exceeds expansion limits")
            if any(e.flag_bits & 1 or "vbaproject" in e.filename.lower() for e in entries):
                raise ValueError("Encrypted or macro-enabled Office files are unsupported")
    return fmt, io.BytesIO(data)


def word(operation, source, spec):
    from docx import Document

    doc = Document(source) if source else Document()
    if operation == "read":
        values = [p.text for p in doc.paragraphs]
        values.extend("\t".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
        content = "\n".join(values)
        return {
            "text": content[:MAX_TEXT],
            "truncated": len(content) > MAX_TEXT,
            "body_paragraphs": [
                {"index": i, "text": p.text[:1000]} for i, p in enumerate(doc.paragraphs[:100])
            ],
        }, None
    if operation == "create":
        fields(spec, ("title", "paragraphs"))
        if spec.get("title"):
            doc.add_heading(text(spec["title"]), 0)
        for p in items(spec.get("paragraphs", [])):
            doc.add_paragraph(text(p))
    else:
        fields(spec, ("paragraph_updates", "append_paragraphs"))
        if not spec.get("paragraph_updates") and not spec.get("append_paragraphs"):
            raise ValueError("No Word edits supplied")
        for update in items(spec.get("paragraph_updates", [])):
            fields(update, ("index", "text"))
            doc.paragraphs[index(update["index"], len(doc.paragraphs))].text = text(update["text"])
        for p in items(spec.get("append_paragraphs", [])):
            doc.add_paragraph(text(p))
    out = io.BytesIO()
    doc.save(out)
    # Validate the saved package by loading it again.
    Document(io.BytesIO(out.getvalue()))
    return {"notice": "Body paragraph editing; replaced paragraphs lose inline run styling."}, out


def excel(operation, source, spec):
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple

    book = load_workbook(source, keep_links=False) if source else Workbook()
    if operation == "read":
        sheets = []
        for sheet in book.worksheets[:10]:
            rows = list(
                sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row, 100),
                    max_col=min(sheet.max_column, 30),
                    values_only=True,
                )
            )
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": rows,
                    "truncated": sheet.max_row > 100 or sheet.max_column > 30,
                }
            )
        return {
            "sheets": sheets,
            "sheets_truncated": len(book.worksheets) > 10,
            "formulas_calculated": False,
        }, None
    if operation == "create":
        fields(spec, ("sheet", "rows"))
        sheet = book.active
        sheet.title = text(spec.get("sheet", "Sheet1"))
        for row in items(spec.get("rows", []), 1000):
            sheet.append([cell_value(v) for v in items(row, 50)])
    else:
        fields(spec, ("sheet", "cells"))
        sheet = book[text(spec.get("sheet", book.active.title))]
        updates = items(spec.get("cells", []), 500)
        if not updates:
            raise ValueError("No cell edits supplied")
        for update in updates:
            fields(update, ("cell", "value"))
            address = update.get("cell")
            if not isinstance(address, str) or not re.fullmatch(
                r"[A-Z]{1,3}[1-9][0-9]{0,3}", address
            ):
                raise ValueError("Invalid cell address")
            row, col = coordinate_to_tuple(address)
            if row > 1000 or col > 50:
                raise ValueError("Cell edit exceeds 1000 rows / 50 columns")
            sheet[address] = cell_value(update.get("value"))
    out = io.BytesIO()
    book.save(out)
    check = load_workbook(io.BytesIO(out.getvalue()))
    check.close()
    book.close()
    return {
        "formulas_calculated": False,
        "notice": "Basic cell edits; complex Excel features may not survive openpyxl.",
    }, out


def pdf(operation, source, spec):
    from pypdf import PdfReader, PdfWriter

    if operation == "create":
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        fields(spec, ("title", "paragraphs"))
        pdfmetrics.registerFont(TTFont("MoriKorean", str(FONT)))
        style = ParagraphStyle(
            "Mori", fontName="MoriKorean", fontSize=11, leading=17, wordWrap="CJK"
        )
        out = io.BytesIO()
        flow = []
        for value in ([spec["title"]] if spec.get("title") else []) + items(
            spec.get("paragraphs", [])
        ):
            flow.extend(
                [Paragraph(escape(text(value)).replace("\n", "<br/>"), style), Spacer(1, 10)]
            )
        if not flow:
            raise ValueError("PDF needs at least one paragraph")
        SimpleDocTemplate(out).build(flow)
        PdfReader(io.BytesIO(out.getvalue()))
        return {"notice": "Korean font embedded; rendered layout not visually verified."}, out
    reader = PdfReader(source)
    if reader.is_encrypted or len(reader.pages) > 100:
        raise ValueError("Encrypted PDFs or PDFs over 100 pages are unsupported")
    if operation == "read":
        content = "\n".join((p.extract_text() or "") for p in reader.pages)
        return {
            "pages": len(reader.pages),
            "text": content[:MAX_TEXT],
            "truncated": len(content) > MAX_TEXT,
            "ocr_performed": False,
        }, None
    fields(spec, ("pages", "rotate"))
    pages = items(spec.get("pages", list(range(1, len(reader.pages) + 1))), 100)
    angle = spec.get("rotate", 0)
    if type(angle) is not int or angle not in (0, 90, 180, 270):
        raise ValueError("Rotation must be 0, 90, 180 or 270")
    if not pages:
        raise ValueError("PDF must keep at least one page")
    writer = PdfWriter()
    for page in pages:
        if type(page) is not int:
            raise ValueError("Page numbers must be integers")
        copied = writer.add_page(reader.pages[index(page - 1, len(reader.pages))])
        copied.rotate(angle)
    out = io.BytesIO()
    writer.write(out)
    PdfReader(io.BytesIO(out.getvalue()))
    return {
        "pages": len(pages),
        "notice": "Page selection/order/rotation only; no PDF text edits.",
    }, out


def execute(args, root):
    fields(args, ("operation", "format", "file_id", "spec"))
    operation = args.get("operation")
    if operation not in ("create", "read", "edit"):
        raise ValueError("Unknown operation")
    if root.is_symlink():
        raise ValueError("Workspace must not be a symlink")
    root.mkdir(mode=0o700, parents=True, exist_ok=True)
    source = None
    if operation == "create":
        if "file_id" in args:
            raise ValueError("Create cannot overwrite a file_id")
        fmt = args.get("format")
    else:
        if "format" in args:
            raise ValueError("Read/edit infer the format from file_id")
        fmt, source = read_input(root, args.get("file_id"))
    spec = args.get("spec", {})
    if not isinstance(spec, dict) or (operation == "read" and spec):
        raise ValueError("Invalid spec")
    handlers = {"docx": word, "xlsx": excel, "pdf": pdf}
    if fmt not in handlers:
        raise ValueError("Only docx, xlsx and pdf are supported")
    details, out = handlers[fmt](operation, source, spec)
    result = {
        "success": True,
        "format": fmt,
        "data": details,
        "external_data_untrusted": True,
        "render_verified": False,
    }
    if out is None:
        result["file_id"] = args["file_id"]
        # Bound previews across multiple cells, tables and sheets as well as paragraphs.
        if len(json.dumps(result, ensure_ascii=False, default=str).encode()) > 200_000:
            raise ValueError("Preview exceeds 200 KB; use a smaller input document")
        return result
    data = out.getvalue()
    if len(data) > MAX_FILE:
        raise ValueError("Output exceeds 10 MB")
    file_id = uuid.uuid4().hex + "." + fmt
    fd = os.open(root / file_id, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "wb") as stream:
        stream.write(data)
    result.update(file_id=file_id, bytes=len(data), source_file_id=args.get("file_id"))
    return result


if __name__ == "__main__":
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
        resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
        raw = sys.stdin.buffer.read(256_001)
        if len(raw) > 256_000:
            raise ValueError("Request exceeds 256 KB")
        home = Path(os.environ.get("HERMES_HOME", "/opt/data"))
        result = execute(json.loads(raw), home / "documents")
    except Exception as exc:
        result = {"success": False, "error": f"Document operation failed ({type(exc).__name__})"}
    print(json.dumps(result, ensure_ascii=False, default=str))
