import importlib.util
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[2]


def load(name, filename):
    spec = importlib.util.spec_from_file_location(
        name, ROOT / "hermes/plugins/mori_documents" / filename
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


worker = load("document_worker", "worker.py")
plugin = load("document_plugin", "__init__.py")


class DocumentsTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name) / "documents"

    def run_doc(self, operation, **kwargs):
        return worker.execute({"operation": operation, **kwargs}, self.root)

    def create(self, fmt, **spec):
        return self.run_doc("create", format=fmt, spec=spec)["file_id"]

    def test_word_edit_preserves_original_and_unmodified_format(self):
        file_id = self.create("docx", paragraphs=["기존 일정", "계속 유지"])
        before = (self.root / file_id).read_bytes()
        changed = self.run_doc(
            "edit",
            file_id=file_id,
            spec={
                "paragraph_updates": [{"index": 0, "text": "새 일정"}],
                "append_paragraphs": ["주차: 지하 2층"],
            },
        )
        self.assertNotEqual(changed["file_id"], file_id)
        self.assertEqual(before, (self.root / file_id).read_bytes())
        data = self.run_doc("read", file_id=changed["file_id"])["data"]
        self.assertIn("새 일정\n계속 유지\n주차: 지하 2층", data["text"])
        self.assertEqual(os.stat(self.root / changed["file_id"]).st_mode & 0o777, 0o600)

    def test_existing_word_table_survives_paragraph_edit(self):
        file_id = self.create("docx", paragraphs=["before"])
        doc = Document(self.root / file_id)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "keep table"
        doc.save(self.root / file_id)
        out = self.run_doc(
            "edit", file_id=file_id, spec={"paragraph_updates": [{"index": 0, "text": "after"}]}
        )
        self.assertEqual(
            Document(self.root / out["file_id"]).tables[0].cell(0, 0).text, "keep table"
        )

    def test_xlsx_cell_edit_and_formula_roundtrip(self):
        file_id = self.create("xlsx", sheet="여행", rows=[["항목", "비용"], ["식사", 10000]])
        before = (self.root / file_id).read_bytes()
        out = self.run_doc(
            "edit",
            file_id=file_id,
            spec={
                "sheet": "여행",
                "cells": [{"cell": "B2", "value": 15000}, {"cell": "B3", "value": "=SUM(B2:B2)"}],
            },
        )
        book = load_workbook(self.root / out["file_id"])
        self.assertEqual(book["여행"]["B2"].value, 15000)
        self.assertEqual(book["여행"]["B3"].value, "=SUM(B2:B2)")
        self.assertFalse(out["data"]["formulas_calculated"])
        self.assertEqual(before, (self.root / file_id).read_bytes())
        self.assertEqual(
            self.run_doc("read", file_id=out["file_id"])["data"]["sheets"][0]["name"], "여행"
        )
        book.close()

    def test_pdf_page_selection_rotation_preserves_source(self):
        self.root.mkdir()
        file_id = "a" * 32 + ".pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=400)
        writer.add_blank_page(width=500, height=600)
        writer.write(self.root / file_id)
        before = (self.root / file_id).read_bytes()
        out = self.run_doc("edit", file_id=file_id, spec={"pages": [2], "rotate": 90})
        reader = PdfReader(self.root / out["file_id"])
        self.assertEqual(len(reader.pages), 1)
        self.assertEqual(reader.pages[0].mediabox.width, 500)
        self.assertEqual(reader.pages[0].rotation, 90)
        self.assertEqual(before, (self.root / file_id).read_bytes())

    def test_unsupported_format_and_operations(self):
        for fmt in ["hwp", "hwpx", "doc", "xls", "xlsm"]:
            with self.subTest(fmt=fmt), self.assertRaises(ValueError):
                self.create(fmt)
        with self.assertRaises(ValueError):
            self.run_doc("execute", spec={"code": "print(1)"})

    def test_paths_and_symlink_inputs_rejected(self):
        for file_id in ["../secret.docx", "/opt/data/config.yaml", "https://example.com/a.docx"]:
            with self.subTest(file_id=file_id), self.assertRaises(ValueError):
                self.run_doc("read", file_id=file_id)
        target = self.root.parent / "private"
        target.write_text("secret")
        (self.root / ("a" * 32 + ".docx")).symlink_to(target)
        with self.assertRaises(OSError):
            self.run_doc("read", file_id="a" * 32 + ".docx")

    def test_workspace_symlink_rejected(self):
        self.root.symlink_to(self.root.parent, target_is_directory=True)
        with self.assertRaises(ValueError):
            self.create("docx", paragraphs=["test"])

    def test_create_cannot_overwrite_input(self):
        with self.assertRaises(ValueError):
            self.run_doc("create", format="docx", file_id="a" * 32 + ".docx")

    def test_unknown_spec_and_invalid_paragraph_do_not_write_output(self):
        file_id = self.create("docx", paragraphs=["keep"])
        for spec in [{"shell": "anything"}, {"paragraph_updates": [{"index": 99, "text": "bad"}]}]:
            with self.assertRaises(ValueError):
                self.run_doc("edit", file_id=file_id, spec=spec)
        self.assertEqual(len(list(self.root.iterdir())), 1)

    def test_cell_bounds_and_invalid_values(self):
        file_id = self.create("xlsx")
        for cell in ["A0", "AY1", "A1001", "A1:B3"]:
            with self.subTest(cell=cell), self.assertRaises(ValueError):
                self.run_doc("edit", file_id=file_id, spec={"cells": [{"cell": cell, "value": 3}]})
        with self.assertRaises(ValueError):
            self.create("xlsx", rows=[[float("nan")]])

    def test_zip_expansion_and_macros_rejected(self):
        self.root.mkdir()
        file_id = "a" * 32 + ".xlsx"
        for name, data in [("xl/vbaProject.bin", b"anything"), ("large", b"0" * 1024)]:
            with zipfile.ZipFile(self.root / file_id, "w") as z:
                z.writestr(name, data)
            if name == "large":
                with patch.object(worker, "MAX_FILE", 10), self.assertRaises(ValueError):
                    self.run_doc("read", file_id=file_id)
            else:
                with self.assertRaises(ValueError):
                    self.run_doc("read", file_id=file_id)

    def test_expanded_archive_limit_rejected_before_parser(self):
        self.root.mkdir()
        file_id = "d" * 32 + ".docx"
        with zipfile.ZipFile(self.root / file_id, "w") as archive:
            archive.writestr("small", "data")
        inflated = zipfile.ZipInfo("large.xml")
        inflated.file_size = 50_000_001
        with patch.object(zipfile.ZipFile, "infolist", return_value=[inflated]):
            with self.assertRaisesRegex(ValueError, "expansion"):
                self.run_doc("read", file_id=file_id)

    def test_encrypted_pdf_rejected(self):
        self.root.mkdir()
        file_id = "b" * 32 + ".pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt("secret")
        writer.write(self.root / file_id)
        with self.assertRaises(ValueError):
            self.run_doc("read", file_id=file_id)

    def test_pdf_invalid_pages_and_angle(self):
        self.root.mkdir()
        file_id = "c" * 32 + ".pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(self.root / file_id)
        for spec in [{"pages": [0]}, {"pages": []}, {"rotate": 45}, {"pages": [True]}]:
            with self.assertRaises(ValueError):
                self.run_doc("edit", file_id=file_id, spec=spec)

    def test_plugin_registration_and_no_credential_forwarding(self):
        class Context:
            def register_tool(self, **kwargs):
                self.tool = kwargs

        ctx = Context()
        plugin.register(ctx)
        self.assertEqual(ctx.tool["name"], "mori_document")
        with patch.dict(os.environ, {"LLAMA_API_KEY": "must-not-forward"}):
            with patch.object(plugin.subprocess, "run") as run:
                run.return_value.returncode = 0
                run.return_value.stdout = '{"success":true}'
                self.assertTrue(json.loads(plugin.document({"operation": "read"}))["success"])
                self.assertNotIn("LLAMA_API_KEY", run.call_args.kwargs["env"])
                self.assertEqual(run.call_args.kwargs["timeout"], 30)

    def test_plugin_timeout_not_success(self):
        with patch.object(
            plugin.subprocess, "run", side_effect=plugin.subprocess.TimeoutExpired("worker", 30)
        ):
            self.assertFalse(json.loads(plugin.document({"operation": "read"}))["success"])


if __name__ == "__main__":
    unittest.main()
